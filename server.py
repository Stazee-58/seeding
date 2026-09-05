"""
server.py — Flask Server & REST APIs cho Hệ thống FB Tracking No-API Pro.
- Quản lý xác thực người dùng (Login / Register / Pending / Logout)
- Quản trị Admin kích hoạt tài khoản (/admin, toggle-active, delete-user)
- API Quét tương tác bài viết không cần Meta Graph API (/api/start-tracking)
- API Kiểm tra Cookie Facebook (/api/check-cookie)
- API Xuất báo cáo Excel 4 sheets (/download/report/<filename>)
"""

import os
import re
import shutil
import sys
import tempfile
from typing import Any, Dict, List

from flask import (
    Flask,
    jsonify,
    make_response,
    redirect,
    render_template,
    request,
    send_file,
    url_for,
)
import openpyxl

import database
from matcher import get_stranger_interactions, match_all
from no_api_tracker import (
    NoApiFacebookTracker,
    extract_post_id_from_url,
    validate_facebook_cookie,
)
from report import generate_excel_report

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "fb_tracking_no_api_super_secret_2026")

# Khởi tạo database
database.init_db()

REPORTS_DIR = os.path.join(tempfile.gettempdir(), "fb_tracking_reports")
os.makedirs(REPORTS_DIR, exist_ok=True)


# ---------------------------------------------------------------------------
# Helper: Kiểm tra xác thực User qua Session Cookie
# ---------------------------------------------------------------------------
def get_current_user():
    session_id = request.cookies.get("tracking_session_id")
    if not session_id:
        return None
    user = database.get_session_user(session_id)
    if not user and os.environ.get("VERCEL"):
        # Khôi phục phiên Admin trên Vercel container mới
        user = database.get_user_by_email(database.ADMIN_EMAIL)
    return user


def parse_members_excel(filepath_or_stream) -> List[Dict[str, Any]]:
    """Đọc file Excel danh sách thành viên (cột: Họ và tên, Facebook ID, Tên hiển thị)."""
    members = []
    try:
        wb = openpyxl.load_workbook(filepath_or_stream, data_only=True)
        ws = wb.active
        headers = []
        for cell in ws[1]:
            val = str(cell.value or "").strip().lower()
            headers.append(val)

        # Xác định chỉ số cột
        col_name = 0
        col_id = 1
        col_alias = 2

        for idx, h in enumerate(headers):
            if "họ" in h or "tên" in h or "name" in h:
                col_name = idx
            elif "id" in h or "uid" in h or "link" in h:
                col_id = idx
            elif "hiển thị" in h or "alias" in h or "phụ" in h:
                col_alias = idx

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue
            name = str(row[col_name] or "").strip()
            if not name:
                continue

            fb_id = str(row[col_id] or "").strip() if len(row) > col_id else ""
            if fb_id.endswith(".0"):
                fb_id = fb_id[:-2]

            alias = str(row[col_alias] or "").strip() if len(row) > col_alias else ""

            members.append({
                "ho_ten": name,
                "facebook_id": fb_id,
                "ten_hien_thi": alias,
            })
    except Exception as e:
        print(f"[-] Lỗi đọc file Excel thành viên: {e}")

    return members


def parse_posts_excel(filepath_or_stream) -> List[str]:
    """Đọc file Excel danh sách link bài viết."""
    urls = []
    try:
        wb = openpyxl.load_workbook(filepath_or_stream, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue
            for cell_val in row:
                val = str(cell_val or "").strip()
                if "facebook.com" in val or "fb.watch" in val:
                    urls.append(val)
                    break
    except Exception as e:
        print(f"[-] Lỗi đọc file Excel bài viết: {e}")
    return urls


# ---------------------------------------------------------------------------
# Routes: Xác thực (Auth)
# ---------------------------------------------------------------------------
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "GET":
        user = get_current_user()
        if user:
            return redirect(url_for("index"))
        return render_template("login.html")

    email = request.form.get("email", "").strip().lower()
    password = request.form.get("password", "")

    user = database.get_user_by_email(email)
    if not user or not database.verify_password(password, user["password_hash"]):
        return render_template("login.html", error="Email hoặc mật khẩu không chính xác!", email=email)

    if user["is_active"] == 0 and user["role"] != "admin":
        session_id = database.create_session(user["id"])
        resp = make_response(redirect(url_for("pending")))
        resp.set_cookie("tracking_session_id", session_id, max_age=7*86400, httponly=True)
        return resp

    session_id = database.create_session(user["id"])
    resp = make_response(redirect(url_for("index")))
    resp.set_cookie("tracking_session_id", session_id, max_age=7*86400, httponly=True)
    return resp


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "GET":
        return render_template("register.html")

    full_name = request.form.get("full_name", "").strip()
    email = request.form.get("email", "").strip().lower()
    phone = request.form.get("phone", "").strip()
    password = request.form.get("password", "")
    confirm_password = request.form.get("confirm_password", "")

    if not full_name or not email or not password:
        return render_template("register.html", error="Vui lòng điền đầy đủ các thông tin bắt buộc!")

    if password != confirm_password:
        return render_template("register.html", error="Mật khẩu xác nhận không khớp!")

    if len(password) < 6:
        return render_template("register.html", error="Mật khẩu phải có ít nhất 6 ký tự!")

    existing = database.get_user_by_email(email)
    if existing:
        return render_template("register.html", error="Email này đã được đăng ký trên hệ thống!")

    database.create_user(email=email, password=password, full_name=full_name, phone=phone, role="user", is_active=0)
    return render_template("register.html", success=True)


@app.route("/logout")
def logout():
    session_id = request.cookies.get("tracking_session_id")
    if session_id:
        database.delete_session(session_id)
    resp = make_response(redirect(url_for("login")))
    resp.delete_cookie("tracking_session_id")
    return resp


@app.route("/pending")
def pending():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["is_active"] == 1 or user["role"] == "admin":
        return redirect(url_for("index"))
    return render_template("pending.html", current_user=user)


# ---------------------------------------------------------------------------
# Routes: Quản trị Admin
# ---------------------------------------------------------------------------
@app.route("/admin")
def admin():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "admin":
        return redirect(url_for("index"))

    users = database.get_all_users()
    return render_template("admin.html", current_user=user, users=users)


@app.route("/api/admin/toggle-active", methods=["POST"])
def admin_toggle_active():
    user = get_current_user()
    if not user or user["role"] != "admin":
        return jsonify({"success": False, "error": "Không có quyền quản trị"}), 403

    data = request.get_json() or {}
    user_id = data.get("user_id")
    is_active = data.get("is_active", 0)

    if not user_id:
        return jsonify({"success": False, "error": "Thiếu user_id"}), 400

    success = database.set_user_active(int(user_id), int(is_active))
    return jsonify({"success": success})


@app.route("/api/admin/delete-user", methods=["POST"])
def admin_delete_user():
    user = get_current_user()
    if not user or user["role"] != "admin":
        return jsonify({"success": False, "error": "Không có quyền quản trị"}), 403

    data = request.get_json() or {}
    user_id = data.get("user_id")
    if not user_id:
        return jsonify({"success": False, "error": "Thiếu user_id"}), 400

    success = database.delete_user(int(user_id))
    return jsonify({"success": success})


# ---------------------------------------------------------------------------
# Routes: Dashboard Chính & Quét Không API
# ---------------------------------------------------------------------------
@app.route("/")
def index():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["is_active"] == 0 and user["role"] != "admin":
        return redirect(url_for("pending"))

    return render_template("index.html", current_user=user)


@app.route("/api/check-cookie", methods=["POST"])
def check_cookie():
    data = request.get_json() or {}
    cookie_str = data.get("cookie", "")
    res = validate_facebook_cookie(cookie_str)
    return jsonify(res)


def sanitize_for_json(obj: Any) -> Any:
    """Đảm bảo mọi chuỗi text không chứa ký tự surrogate lỗi gây crash UTF-8."""
    if isinstance(obj, str):
        try:
            return obj.encode('utf-16', 'surrogatepass').decode('utf-16').encode('utf-8', 'ignore').decode('utf-8')
        except Exception:
            try:
                return obj.encode('utf-8', 'ignore').decode('utf-8')
            except Exception:
                return "".join(c for c in obj if ord(c) < 0xD800 or ord(c) > 0xDFFF)
    elif isinstance(obj, dict):
        return {k: sanitize_for_json(v) for k, v in obj.items()}
    elif isinstance(obj, list):
        return [sanitize_for_json(x) for x in obj]
    return obj


def _handle_start_tracking():
    user = get_current_user()
    if not user:
        return jsonify({"success": False, "error": "Vui lòng đăng nhập"}), 401
    if user["is_active"] == 0 and user["role"] != "admin":
        return jsonify({"success": False, "error": "Tài khoản đang chờ duyệt"}), 403

    cookie_str = request.form.get("cookie", "").strip()
    mode = request.form.get("mode", "page").strip()
    page_url = request.form.get("page_url", "").strip()
    max_posts_str = request.form.get("max_posts", "0").strip()
    post_urls_text = request.form.get("post_urls", "").strip()
    check_likes = request.form.get("check_likes") == "1"
    check_comments = request.form.get("check_comments") == "1"
    check_shares = request.form.get("check_shares") == "1"

    try:
        max_posts = int(max_posts_str)
    except ValueError:
        max_posts = 0

    logs_list = []
    def log_cb(msg: str, lvl: str = "info"):
        logs_list.append({"msg": msg, "level": lvl})

    tracker = NoApiFacebookTracker(cookie_str=cookie_str, log_callback=log_cb)

    target_posts = []

    # -------------------------------------------------------------------------
    # CHẾ ĐỘ 1: QUÉT TOÀN BỘ BÀI VIẾT TRÊN FANPAGE
    # -------------------------------------------------------------------------
    if mode == "page" or (page_url and not post_urls_text and "post_file" not in request.files):
        if not page_url:
            return jsonify({"success": False, "error": "Vui lòng nhập đường dẫn Fanpage hoặc ID Trang!"})

        target_posts = tracker.fetch_page_posts(page_url, limit=max_posts)
        if not target_posts:
            return jsonify({
                "success": False,
                "error": f"Không tìm thấy bài viết nào trên Trang '{page_url}'. Hãy kiểm tra lại link hoặc cookie đăng nhập!",
                "logs": logs_list,
            })

    # -------------------------------------------------------------------------
    # CHẾ ĐỘ 2: QUÉT THEO DANH SÁCH BÀI VIẾT HOẶC FILE
    # -------------------------------------------------------------------------
    else:
        post_urls = []
        if post_urls_text:
            for line in post_urls_text.splitlines():
                line = line.strip()
                if line and ("facebook.com" in line or "fb.watch" in line):
                    post_urls.append(line)

        if "post_file" in request.files:
            p_file = request.files["post_file"]
            if p_file and p_file.filename:
                file_urls = parse_posts_excel(p_file)
                post_urls.extend(file_urls)

        raw_urls = list(dict.fromkeys(post_urls))
        if not raw_urls:
            return jsonify({"success": False, "error": "Không tìm thấy link bài viết hợp lệ nào!"})

        # Kiểm tra nếu trong danh sách có link Fanpage -> tự động cào bài của trang đó
        for u in raw_urls:
            if extract_post_id_from_url(u):
                target_posts.append(u)
            else:
                page_posts = tracker.fetch_page_posts(u, limit=max_posts if max_posts > 0 else 15)
                if page_posts:
                    target_posts.extend(page_posts)
                else:
                    target_posts.append(u)

    tracker.log(f"🎯 Bắt đầu phân tích tương tác trên tổng cộng {len(target_posts)} bài viết...", "info")

    # Lấy danh sách thành viên (nếu người dùng tải file lên để đối soát)
    members = []
    has_members = False
    if "member_file" in request.files:
        m_file = request.files["member_file"]
        if m_file and m_file.filename:
            members = parse_members_excel(m_file)
            if members:
                has_members = True
                tracker.log(f"👥 Đã nạp {len(members)} thành viên từ file Excel để đối soát.", "info")

    if not has_members:
        tracker.log("ℹ️ Chế độ: Quét danh sách người tương tác toàn trang (Không dùng danh sách thành viên).", "info")

    # Tiến hành cào dữ liệu thực tế 100% không cần API
    raw_posts_data = tracker.track_posts(
        post_urls=target_posts,
        check_likes=check_likes,
        check_comments=check_comments,
        check_shares=check_shares,
    )

    # Chuẩn hoá sang dict collected và posts_details
    collected = {}
    posts_details = []

    for p in raw_posts_data:
        pid = p["post_id"]
        mo_ta = p.get("mo_ta", f"Bài ID {pid}")
        p_url = p.get("post_url", "")
        reacts = p.get("reactions", [])
        cmts = p.get("comments", [])
        shrs = p.get("shares", [])

        collected[pid] = {
            "mo_ta": mo_ta,
            "permalink_url": p_url,
            "comments": cmts,
            "reactions": reacts,
            "shares": shrs,
            "likes_count": p.get("likes_count", len(reacts)),
            "comments_count": p.get("comments_count", len(cmts)),
            "shares_count": p.get("shares_count", len(shrs)),
            "total_reactions": p.get("total_reactions", len(reacts)),
            "total_shares": p.get("total_shares", len(shrs)),
        }

        posts_details.append({
            "post_id": pid,
            "mo_ta": mo_ta,
            "permalink_url": p_url,
            "likes_count": p.get("likes_count", len(reacts)),
            "comments_count": p.get("comments_count", len(cmts)),
            "shares_count": p.get("shares_count", len(shrs)),
            "total_reactions": p.get("total_reactions", len(reacts)),
            "total_shares": p.get("total_shares", len(shrs)),
            "reactions": reacts,
            "comments": cmts,
            "shares": shrs,
        })

    # =========================================================================
    # TỔNG HỢP DANH SÁCH TẤT CẢ NGƯỜI DÙNG TƯƠNG TÁC TOÀN TRANG (CHECK TÊN RA)
    # =========================================================================
    user_engagement = {}
    for p in raw_posts_data:
        pid = str(p["post_id"])
        p_url = p.get("post_url", "")
        mo_ta = p.get("mo_ta") or f"Bài ID {pid}"

        # 1. Thả tim / Like
        for r in p.get("reactions", []):
            u_name = r.get("name", "").strip()
            u_id = r.get("id", "").strip()
            if not u_name:
                continue
            key = u_id if u_id else u_name.lower()
            if key not in user_engagement:
                user_engagement[key] = {
                    "name": u_name,
                    "user_id": u_id,
                    "profile_url": r.get("profile_url") or (f"https://facebook.com/{u_id}" if u_id else ""),
                    "likes_count": 0,
                    "comments_count": 0,
                    "shares_count": 0,
                    "total_interactions": 0,
                    "liked_posts": [],
                    "commented_posts": [],
                    "shared_posts": [],
                }
            user_engagement[key]["likes_count"] += 1
            user_engagement[key]["total_interactions"] += 1
            if pid not in [x["post_id"] for x in user_engagement[key]["liked_posts"]]:
                user_engagement[key]["liked_posts"].append({
                    "post_id": pid,
                    "type": r.get("type", "LIKE"),
                    "mo_ta": mo_ta,
                })

        # 2. Bình luận
        for c in p.get("comments", []):
            u_name = (c.get("from_name") or c.get("name") or "").strip()
            u_id = (c.get("from_id") or c.get("id") or "").strip()
            if not u_name:
                continue
            key = u_id if u_id else u_name.lower()
            if key not in user_engagement:
                user_engagement[key] = {
                    "name": u_name,
                    "user_id": u_id,
                    "profile_url": f"https://facebook.com/{u_id}" if u_id else "",
                    "likes_count": 0,
                    "comments_count": 0,
                    "shares_count": 0,
                    "total_interactions": 0,
                    "liked_posts": [],
                    "commented_posts": [],
                    "shared_posts": [],
                }
            user_engagement[key]["comments_count"] += 1
            user_engagement[key]["total_interactions"] += 1
            user_engagement[key]["commented_posts"].append({
                "post_id": pid,
                "message": c.get("message") or "",
                "created_time": c.get("created_time") or "",
                "mo_ta": mo_ta,
            })

        # 3. Chia sẻ
        for s in p.get("shares", []):
            u_name = s.get("name", "").strip()
            u_id = s.get("id", "").strip()
            if not u_name:
                continue
            key = u_id if u_id else u_name.lower()
            if key not in user_engagement:
                user_engagement[key] = {
                    "name": u_name,
                    "user_id": u_id,
                    "profile_url": f"https://facebook.com/{u_id}" if u_id else "",
                    "likes_count": 0,
                    "comments_count": 0,
                    "shares_count": 0,
                    "total_interactions": 0,
                    "liked_posts": [],
                    "commented_posts": [],
                    "shared_posts": [],
                }
            user_engagement[key]["shares_count"] += 1
            user_engagement[key]["total_interactions"] += 1
            if pid not in user_engagement[key]["shared_posts"]:
                user_engagement[key]["shared_posts"].append(pid)

    # Tính tỷ lệ tham gia (%) và gom thành danh sách xếp hạng
    total_posts_count = len(collected)
    users_directory = []
    for key, u in user_engagement.items():
        distinct_posts = set(
            [x["post_id"] for x in u["liked_posts"]] +
            [x["post_id"] for x in u["commented_posts"]] +
            u["shared_posts"]
        )
        p_cnt = len(distinct_posts)
        rate = round((p_cnt / total_posts_count) * 100, 1) if total_posts_count > 0 else 0.0
        u["distinct_posts_count"] = p_cnt
        u["participation_rate"] = rate
        users_directory.append(u)

    # Sắp xếp người tương tác nhiều nhất lên đầu
    users_directory.sort(key=lambda x: (x["total_interactions"], x["distinct_posts_count"]), reverse=True)

    # Đối soát thành viên nếu có
    match_results = []
    strangers = []
    if has_members and members:
        match_results = match_all(members=members, collected=collected)
        strangers = get_stranger_interactions(members=members, collected=collected)

    # Thống kê bài viết
    post_summary = []
    for pid, pdata in collected.items():
        likes_cnt = pdata.get("likes_count", len(pdata.get("reactions", [])))
        cmts_cnt = pdata.get("comments_count", len(pdata.get("comments", [])))
        shares_cnt = pdata.get("shares_count", len(pdata.get("shares", [])))

        total_mems = len(members)
        if total_mems > 0 and match_results:
            done_cnt = sum(
                1 for r in match_results 
                if r["post_id"] == pid and r["trang_thai"] != "chua_tuong_tac"
            )
            rate = round((done_cnt / total_mems) * 100, 1)
        else:
            rate = 0.0

        post_summary.append({
            "post_id": pid,
            "mo_ta": pdata.get("mo_ta", pid),
            "permalink_url": pdata.get("permalink_url", ""),
            "likes_count": likes_cnt,
            "comments_count": cmts_cnt,
            "shares_count": shares_cnt,
            "completion_rate": rate,
        })

    # Xuất file Excel báo cáo
    report_file_path = generate_excel_report(
        match_results=match_results,
        stranger_results=strangers,
        collected=collected,
        has_members=has_members,
        users_directory=users_directory,
        output_dir=REPORTS_DIR,
    )
    report_filename = os.path.basename(report_file_path)

    # Lưu lịch sử
    campaign_label = f"Quét Trang ({len(collected)} bài, {len(users_directory)} người)" if mode == "page" else f"Quét {len(collected)} bài viết"
    database.add_tracking_history(
        user_id=user["id"],
        campaign_name=campaign_label,
        post_count=len(collected),
        member_count=len(users_directory),
        report_path=report_file_path,
    )

    response_payload = {
        "success": True,
        "has_members": has_members,
        "total_posts": len(collected),
        "total_users_interacted": len(users_directory),
        "total_members": len(members),
        "logs": logs_list,
        "users_directory": users_directory,
        "posts_details": posts_details,
        "match_results": match_results,
        "strangers": strangers,
        "post_summary": post_summary,
        "report_download_url": f"/download/report/{report_filename}",
    }
    return jsonify(sanitize_for_json(response_payload))


@app.route("/api/start-tracking", methods=["POST"])
def start_tracking():
    try:
        return _handle_start_tracking()
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({
            "success": False,
            "error": f"Lỗi hệ thống: {str(e)}",
            "logs": [],
        }), 200


@app.route("/download/report/<filename>")
def download_report(filename):
    safe_name = os.path.basename(filename)
    path = os.path.join(REPORTS_DIR, safe_name)
    if os.path.exists(path):
        return send_file(path, as_attachment=True, download_name=safe_name)
    return "File báo cáo không tồn tại hoặc đã bị xoá.", 404


@app.route("/download/sample/members")
def download_sample_members():
    path = os.path.join(os.path.dirname(__file__), "members.xlsx")
    if os.path.exists(path):
        return send_file(path, as_attachment=True, download_name="members_sample.xlsx")
    return "Không tìm thấy file mẫu.", 404


@app.route("/download/sample/posts")
def download_sample_posts():
    path = os.path.join(os.path.dirname(__file__), "posts.xlsx")
    if os.path.exists(path):
        return send_file(path, as_attachment=True, download_name="posts_sample.xlsx")
    return "Không tìm thấy file mẫu.", 404


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5050))
    print(f"[*] FB Tracking No-API Server dang chay tai: http://localhost:{port}")
    app.run(host="0.0.0.0", port=port, debug=False)
